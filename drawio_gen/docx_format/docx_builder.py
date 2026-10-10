import re

from io import BytesIO

from docx import Document
from docx.shared import Pt, RGBColor
from pygments import highlight
from pygments.lexers import get_lexer_by_name, guess_lexer
from pygments.formatters import get_formatter_by_name
from pygments.styles import get_style_by_name
from pygments.token import Token
from pygments.util import ClassNotFound


_STYLE_CACHE: dict | None = None


def _get_token_styles() -> dict:
    global _STYLE_CACHE
    if _STYLE_CACHE is None:
        style_cls = get_style_by_name("default")
        _STYLE_CACHE = dict(style_cls.styles)
    return _STYLE_CACHE


def _parse_style_def(style_def: str) -> tuple[str | None, bool, bool]:
    color: str | None = None
    bold = False
    italic = False
    if not style_def:
        return None, False, False
    tokens = style_def.strip().split()
    for tok in tokens:
        if tok.startswith("#"):
            color = tok.lstrip("#")
        elif tok == "bold":
            bold = True
        elif tok == "nobold":
            bold = False
        elif tok == "italic":
            italic = True
        elif tok == "noitalic":
            italic = False
    return color, bold, italic


def _rgb(hex_str: str | None) -> RGBColor | None:
    if not hex_str:
        return None
    try:
        return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))
    except (ValueError, IndexError):
        return None


def build_docx(code: str, language: str = "auto") -> BytesIO:
    if language and language != "auto":
        try:
            lexer = get_lexer_by_name(language, stripall=False)
        except ClassNotFound:
            lexer = guess_lexer(code)
    else:
        lexer = guess_lexer(code)

    token_styles = _get_token_styles()

    doc = Document()

    section = doc.sections[0]
    section.top_margin = Pt(36)
    section.bottom_margin = Pt(36)
    section.left_margin = Pt(36)
    section.right_margin = Pt(36)

    style = doc.styles["Normal"]
    style.font.name = "Consolas"
    style.font.size = Pt(10)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.15

    tokens = list(lexer.get_tokens(code))
    paragraph = None

    for ttype, value in tokens:
        if not value:
            continue
        parts = value.split("\n")
        for i, part in enumerate(parts):
            if i > 0:
                paragraph = None
            if not part:
                continue
            if paragraph is None:
                paragraph = doc.add_paragraph()
                pf = paragraph.paragraph_format
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                pf.line_spacing = 1.15

            style_def = token_styles.get(ttype, "")
            color_hex, bold, italic = _parse_style_def(style_def)

            run = paragraph.add_run(part)
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            run.font.bold = bold
            run.font.italic = italic
            rgb = _rgb(color_hex)
            if rgb is not None:
                run.font.color.rgb = rgb

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def highlight_html(code: str, language: str = "auto") -> tuple[str, str]:
    if language and language != "auto":
        try:
            lexer = get_lexer_by_name(language, stripall=False)
        except ClassNotFound:
            lexer = guess_lexer(code)
    else:
        lexer = guess_lexer(code)

    html = highlight(code, lexer, get_formatter_by_name("html", nowrap=True, noclasses=True, style="default"))
    return html, lexer.name


__all__ = ["build_docx", "highlight_html"]